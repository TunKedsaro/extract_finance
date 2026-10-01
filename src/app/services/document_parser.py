import base64
import io
from pathlib import Path
from typing import Final

import fitz
import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..core.exceptions import DocumentProcessingError, UnsupportedDocumentTypeError


PDF_MIME_TYPE: Final = "application/pdf"
XLSX_MIME_TYPE: Final = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
DOCX_MIME_TYPE: Final = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
MIME_TYPES: Final = {
    ".pdf": PDF_MIME_TYPE,
    ".xlsx": XLSX_MIME_TYPE,
    ".docx": DOCX_MIME_TYPE,
}


class DocumentParser:
    """Convert supported documents into Anthropic message content blocks."""

    def mime_type_for(self, file_path: Path) -> str:
        try:
            return MIME_TYPES[file_path.suffix.lower()]
        except KeyError as exc:
            raise UnsupportedDocumentTypeError() from exc

    def to_content_blocks(self, file_bytes: bytes, mime_type: str) -> list[dict]:
        if mime_type == PDF_MIME_TYPE:
            return self.pdf_to_content_blocks(file_bytes)

        text_converter = {
            XLSX_MIME_TYPE: self.xlsx_to_text,
            DOCX_MIME_TYPE: self.docx_to_text,
        }.get(mime_type)
        if text_converter is None:
            raise UnsupportedDocumentTypeError()

        return [{"type": "text", "text": text_converter(file_bytes)}]

    @staticmethod
    def pdf_to_content_blocks(file_bytes: bytes) -> list[dict]:
        """Route text-heavy PDF pages as text and other pages as JPEG images."""

        blocks: list[dict] = []
        try:
            with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                if pdf.page_count == 0:
                    raise DocumentProcessingError("The PDF contains no pages.")

                for page_number, page in enumerate(pdf, start=1):
                    text = page.get_text("text", sort=True).strip()
                    page_area = page.rect.width * page.rect.height
                    has_large_image = page_area > 0 and any(
                        DocumentParser._image_area_ratio(image, page.rect, page_area)
                        >= 0.20
                        for image in page.get_image_info()
                    )

                    if len(text) >= 200 and not has_large_image:
                        blocks.append(
                            {
                                "type": "text",
                                "text": f"Page {page_number} (text):\n{text}",
                            }
                        )
                        continue

                    image_bytes = page.get_pixmap(
                        matrix=fitz.Matrix(1.5, 1.5),
                        alpha=False,
                    ).tobytes("jpeg", jpg_quality=82)
                    blocks.extend(
                        [
                            {
                                "type": "text",
                                "text": f"Page {page_number} image:",
                            },
                            {
                                "type": "image",
                                "source": {
                                    "type": "base64",
                                    "media_type": "image/jpeg",
                                    "data": base64.b64encode(image_bytes).decode("ascii"),
                                },
                            },
                        ]
                    )
        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentProcessingError("The PDF could not be read.") from exc

        return blocks

    @staticmethod
    def _image_area_ratio(image: dict, page_rect: fitz.Rect, page_area: float) -> float:
        image_rect = fitz.Rect(image["bbox"]) & page_rect
        return image_rect.width * image_rect.height / page_area

    @staticmethod
    def xlsx_to_text(file_bytes: bytes) -> str:
        """Flatten a workbook while retaining sheet and column-position context."""

        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(file_bytes),
                data_only=True,
                read_only=True,
            )
            lines: list[str] = []
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                lines.append(f"=== Sheet: {sheet_name} ===")
                for row in worksheet.iter_rows():
                    cells = [
                        (cell.column - 1, cell.value)
                        for cell in row
                        if cell.value is not None
                    ]
                    if not cells:
                        continue
                    parts = [f"{'  ' * column}{value}" for column, value in cells]
                    lines.append(" | ".join(parts))
                lines.append("")
        except Exception as exc:
            raise DocumentProcessingError("The XLSX workbook could not be read.") from exc
        finally:
            if workbook is not None:
                workbook.close()

        return "\n".join(lines)

    @staticmethod
    def docx_to_text(file_bytes: bytes) -> str:
        """Flatten Word paragraphs and tables in their original document order."""

        try:
            document = Document(io.BytesIO(file_bytes))
            lines: list[str] = []
            table_number = 0

            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if text:
                        lines.append(text)
                    continue

                if isinstance(block, Table):
                    table_number += 1
                    table_lines = DocumentParser._docx_table_lines(block)
                    if not table_lines:
                        continue
                    lines.append(f"=== Table: {table_number} ===")
                    lines.extend(table_lines)
                    lines.append("")

            if not lines:
                raise DocumentProcessingError(
                    "The DOCX document contains no readable text."
                )
        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentProcessingError(
                "The DOCX document could not be read."
            ) from exc

        return "\n".join(lines).strip()

    @staticmethod
    def _docx_table_lines(table: Table) -> list[str]:
        """Represent table rows without losing empty column positions."""

        lines: list[str] = []
        for row_number, row in enumerate(table.rows, start=1):
            values = [
                " ".join(part.strip() for part in cell.text.splitlines() if part.strip())
                for cell in row.cells
            ]
            if any(values):
                lines.append(f"Row {row_number}: {' | '.join(values)}")
        return lines
