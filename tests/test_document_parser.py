import base64
import io
from pathlib import Path

import fitz
import openpyxl
import pytest
from docx import Document

from app.core.exceptions import DocumentProcessingError, UnsupportedDocumentTypeError
from app.services.document_parser import (
    DOCX_MIME_TYPE,
    MIME_TYPES,
    PDF_MIME_TYPE,
    XLSX_MIME_TYPE,
    DocumentParser,
)


def test_supported_mime_types() -> None:
    parser = DocumentParser()

    assert set(MIME_TYPES) == {".pdf", ".xlsx", ".docx"}
    assert parser.mime_type_for(Path("REPORT.PDF")) == PDF_MIME_TYPE
    assert parser.mime_type_for(Path("REPORT.XLSX")) == XLSX_MIME_TYPE
    assert parser.mime_type_for(Path("REPORT.DOCX")) == DOCX_MIME_TYPE

    with pytest.raises(UnsupportedDocumentTypeError):
        parser.mime_type_for(Path("report.csv"))


def test_pdf_routes_text_and_image_pages() -> None:
    image_document = fitz.open()
    image_page = image_document.new_page(width=20, height=20)
    image_page.draw_rect(image_page.rect, color=(0, 0, 0), fill=(1, 1, 1))
    png = image_page.get_pixmap().tobytes("png")
    image_document.close()

    pdf = fitz.open()
    text_page = pdf.new_page()
    text_page.insert_textbox(
        fitz.Rect(30, 40, 580, 750),
        "Financial statement page with selectable text. " * 12,
        fontsize=9,
    )
    scan_page = pdf.new_page()
    scan_page.insert_image(scan_page.rect, stream=png)
    mixed_page = pdf.new_page()
    mixed_page.insert_image(mixed_page.rect, stream=png)
    mixed_page.insert_textbox(
        fitz.Rect(30, 40, 580, 750),
        "Selectable text over a full-page scan. " * 12,
        fontsize=9,
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()

    blocks = DocumentParser.pdf_to_content_blocks(pdf_bytes)

    assert [block["type"] for block in blocks] == [
        "text",
        "text",
        "image",
        "text",
        "image",
    ]
    assert "Page 1 (text)" in blocks[0]["text"]
    assert blocks[1]["text"] == "Page 2 image:"
    assert blocks[3]["text"] == "Page 3 image:"
    assert base64.b64decode(blocks[2]["source"]["data"]).startswith(b"\xff\xd8")
    assert base64.b64decode(blocks[4]["source"]["data"]).startswith(b"\xff\xd8")


def test_invalid_pdf_is_reported_as_document_error() -> None:
    with pytest.raises(DocumentProcessingError):
        DocumentParser.pdf_to_content_blocks(b"not a pdf")


def test_xlsx_preserves_sheet_and_column_position() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Balance Sheet"
    worksheet["A1"] = "Current assets"
    worksheet["B2"] = "Cash"
    worksheet["C2"] = 1250
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    text = DocumentParser.xlsx_to_text(buffer.getvalue())

    assert "=== Sheet: Balance Sheet ===" in text
    assert "Current assets" in text
    assert "  Cash |     1250" in text


def test_docx_preserves_paragraph_and_table_order() -> None:
    document = Document()
    document.add_paragraph("Balance sheet")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[2].text = "Amount"
    table.rows[1].cells[0].text = "Cash"
    table.rows[1].cells[2].text = "1250"
    document.add_paragraph("End of statement")
    buffer = io.BytesIO()
    document.save(buffer)

    blocks = DocumentParser().to_content_blocks(
        buffer.getvalue(),
        DOCX_MIME_TYPE,
    )
    text = blocks[0]["text"]

    assert [block["type"] for block in blocks] == ["text"]
    assert text.index("Balance sheet") < text.index("=== Table: 1 ===")
    assert "Row 1: Item |  | Amount" in text
    assert "Row 2: Cash |  | 1250" in text
    assert text.index("=== Table: 1 ===") < text.index("End of statement")


def test_invalid_docx_is_reported_as_document_error() -> None:
    with pytest.raises(DocumentProcessingError):
        DocumentParser.docx_to_text(b"not a docx")


def test_empty_docx_is_reported_as_document_error() -> None:
    document = Document()
    buffer = io.BytesIO()
    document.save(buffer)

    with pytest.raises(DocumentProcessingError):
        DocumentParser.docx_to_text(buffer.getvalue())
